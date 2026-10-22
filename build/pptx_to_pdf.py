"""Render a .pptx deck to PDF without PowerPoint or LibreOffice.

Usage: python build/pptx_to_pdf.py docs/deck.pptx build/deck.pdf

Draws each slide's background, filled shapes and text boxes at their real
positions, sizes and colours with PyMuPDF. It is an approximation: standard
PDF fonts stand in for the deck's fonts, and shadows, gradients and images
are skipped. Open the .pptx in PowerPoint for the exact look.
"""

import sys

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

EMU_PER_PT = 12700


def pt(v):
    return Emu(v).pt if v is not None else 0.0


def rgb(color, default=(0, 0, 0)):
    try:
        if color is not None and color.type is not None and color.rgb is not None:
            h = str(color.rgb)
            return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except (AttributeError, TypeError, ValueError):
        pass
    return default


def fill_rgb(shape):
    try:
        f = shape.fill
        if f.type == 1:  # MSO_FILL.SOLID
            return rgb(f.fore_color, None)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def slide_background(slide):
    try:
        f = slide.background.fill
        if f.type == 1:
            return rgb(f.fore_color, (1, 1, 1))
    except (AttributeError, TypeError, ValueError):
        pass
    return (1, 1, 1)


ALIGN = {PP_ALIGN.CENTER: fitz.TEXT_ALIGN_CENTER, PP_ALIGN.RIGHT: fitz.TEXT_ALIGN_RIGHT}


def draw_text(page, shape, rect):
    tf = shape.text_frame
    pad = 4
    y = rect.y0 + pad
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs) or para.text
        if not text.strip():
            y += 8
            continue
        run = para.runs[0] if para.runs else None
        size = (run.font.size.pt if run is not None and run.font.size else None) or 18
        bold = bool(run is not None and run.font.bold)
        color = rgb(run.font.color if run is not None else None)
        font = "hebo" if bold else "helv"
        box = fitz.Rect(rect.x0 + pad, y, rect.x1 - pad, rect.y1 + 200)
        align = ALIGN.get(para.alignment, fitz.TEXT_ALIGN_LEFT)
        # insert_textbox returns the unused height (negative if it overflowed).
        left = page.insert_textbox(box, text, fontsize=size, fontname=font, color=color, align=align)
        used = box.height - left if left >= 0 else box.height
        y += used + size * 0.25
        if y > rect.y1 + size:
            break
        # Follow hyperlinks on runs so the white-paper link stays clickable.
        for r in para.runs:
            if r.hyperlink is not None and r.hyperlink.address:
                page.insert_link({"kind": fitz.LINK_URI, "from": box, "uri": r.hyperlink.address})


def render(src, out):
    prs = Presentation(src)
    w, h = pt(prs.slide_width), pt(prs.slide_height)
    doc = fitz.open()
    for slide in prs.slides:
        page = doc.new_page(width=w, height=h)
        page.draw_rect(page.rect, color=None, fill=slide_background(slide))
        for shape in slide.shapes:
            if shape.left is None:
                continue
            rect = fitz.Rect(pt(shape.left), pt(shape.top),
                             pt(shape.left) + pt(shape.width), pt(shape.top) + pt(shape.height))
            fill = fill_rgb(shape)
            if fill is not None:
                page.draw_rect(rect, color=None, fill=fill)
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_text(page, shape, rect)
    doc.save(out)
    print(f"wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    render(sys.argv[1], sys.argv[2])
